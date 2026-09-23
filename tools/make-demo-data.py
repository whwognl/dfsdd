#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
데모(자동화 관제) 데이터 추출 — 실제 워크북에서 '패턴'만 뽑고 개인정보는 한 글자도 담지 않습니다.
  python3 tools/make-demo-data.py <워크북.xlsx>  →  demo-data.js (window.DemoData)
담는 것: 상품·옵션·판매가·매입가·구매처·택배사(빈도), 주문 시간대/요일/수량 분포, 지역(시·도 + 시·군·구까지만),
        일별 매출의 '흔들림'(7일 평균 대비 비율), CS 유형 비율.  담지 않는 것: 이름·전화·주소 상세·계정·계좌·카드·주문번호.
"""
import sys, re, json, datetime, statistics, warnings, collections
warnings.filterwarnings("ignore")
import openpyxl

SRC = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else "demo-data.js"
ANON = "--anon-vendors" in sys.argv          # 구매처 상호를 '구매처 A/B/C…' 로 가명 처리(공개용)
VENDOR_ALIAS = {"pbf": "PBF", "제이비티엠": "JBTM", "웰그린": "웰그린푸드", "웰그린(업푸르트)": "웰그린푸드", "웰그린푸드(업푸르트)": "웰그린푸드"}
def norm_vendor(v):
    v = re.sub(r"\s+", " ", str(v or "")).strip()
    return VENDOR_ALIAS.get(v, VENDOR_ALIAS.get(v.lower(), v))
wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

def rows_of(name):
    ws = wb[name]
    it = ws.iter_rows(values_only=True)
    header = [str(v).replace("\n", " ").strip() if v is not None else "" for v in next(it)]
    return header, it

def num(v):
    try:
        if v is None or v == "": return None
        return float(v)
    except Exception: return None

def clean_name(s):
    s = str(s or "").translate(str.maketrans({"（": "(", "）": ")", "，": ",", "\u3000": " ", "\xa0": " "}))
    s = re.sub(r"\[[^\]]*[\]\}]", " ", s)                                   # [#15] · [특허농법] 같은 대괄호 태그 제거
    s = re.sub(r"재구매폭주|초고당도\s*brix|초고당도|프리미엄|제주직송|전국최저특가|전국최저마진|국내산|산지직송|당일수확|무료배송|엄격선별|해풍맞고자란|시원달큰|남해직송|나주직송|직접조업한|선주직송|쫄깃신선꽁치|톡쏘는|시원아삭|꿀뚝뚝\s*brix|가정용|실속|햇\b|PBF|팡이농장", " ", s, flags=re.I)
    s = re.sub(r"/.*$", " ", s)                                              # '/ 뒤' 부연 제거
    s = re.sub(r",\s*\d+(\.\d+)?\s*(kg|g|개|박스|미|수|팩)?\b.*$", " ", s, flags=re.I)   # 뒤에 붙은 용량
    s = re.sub(r"\(\s*\d+(\.\d+)?\s*(kg|g)\s*\)", " ", s, flags=re.I)
    s = s.replace("_", "")
    s = tidy(s)
    s = s[:24]
    s = re.sub(r"\s*\([^)]*$", "", s).strip()                                # 잘려 열린 괄호
    return tidy(s)

def tidy(s):
    s = re.sub(r"\([^가-힣A-Za-z0-9]*\)", " ", s)                            # 빈 괄호(공백·기호만)
    s = re.sub(r"(\s*[,·]\s*)+", ", ", s)                                     # 쉼표 연속 정리
    s = re.sub(r"[,·]\s*(?=[,·]|$)", " ", s)                                  # 남은 쉼표 조각
    return re.sub(r"\s+", " ", s).strip(" -·/,")

# ---------- 주문 이력 (운송장 전송후) ----------
H, it = rows_of("운송장 전송후")
ix = {h: i for i, h in enumerate(H)}
prod = collections.defaultdict(lambda: {"n": 0, "price": [], "cost": [], "margin": [], "vendor": collections.Counter(), "courier": collections.Counter(), "qty": collections.Counter()})
hours = [0] * 24; wdays = [0] * 7; qtys = collections.Counter(); regions = collections.Counter()
couriers = collections.Counter(); vendors = collections.Counter(); discounts = collections.Counter()
n_rows = 0
for r in it:
    if r is None or len(r) < 20: continue
    name = r[ix["상품명"]]; opt = r[ix["옵션"]]
    if not name: continue
    n_rows += 1
    nm = clean_name(name); op = re.sub(r"\s+", " ", str(opt or "")).strip()[:24]
    if op and op in nm: nm = tidy(nm.replace(op, " "))                  # 상품명에 옵션이 또 들어 있으면 제거
    key = (nm, op)
    p = prod[key]; p["n"] += 1
    sp = num(r[ix["(실)판매가"]]) or num(r[ix["판매가"]]); cp = num(r[ix["매입가"]]); mg = num(r[ix["순마진"]])
    if str(r[ix["거래처"]] or "").strip() == "리뷰" or (sp and cp and cp > sp): p["n"] -= 1; continue
    if sp and sp > 0: p["price"].append(sp)
    if cp and cp > 0: p["cost"].append(cp)
    if mg is not None: p["margin"].append(mg)
    v = norm_vendor(r[ix["거래처"]]); c = str(r[ix["택배사"]] or "").strip()
    if v: p["vendor"][v] += 1; vendors[v] += 1
    if c: p["courier"][c] += 1; couriers[c] += 1
    q = int(num(r[ix["주문수량"]]) or 1); p["qty"][q] += 1; qtys[q] += 1
    d = r[ix["주문일"]]
    if isinstance(d, datetime.datetime): hours[d.hour] += 1; wdays[d.weekday()] += 1
    dc = num(r[ix["할인금액"]]);
    if dc is not None: discounts[int(dc)] += 1
    addr = str(r[ix["주소"]] or "").split()
    if len(addr) >= 2 and re.search(r"(시|군|구)$", addr[1]): regions[addr[0] + " " + addr[1]] += 1
    elif len(addr) >= 1 and addr[0]: regions[addr[0]] += 1

products = []
for (name, opt), p in prod.items():
    if p["n"] < 3 or not p["price"]: continue
    if re.search(r"보스턴백|model|가방|케이스", name, re.I) or len(name) < 2: continue
    price = statistics.median(p["price"]); cost = statistics.median(p["cost"]) if p["cost"] else round(price * 0.72, -1)
    products.append({
        "name": name, "opt": opt, "n": p["n"], "price": int(round(price, -1)), "cost": int(round(cost, -1)),
        "vendor": (p["vendor"].most_common(1) or [("", 0)])[0][0], "courier": (p["courier"].most_common(1) or [("", 0)])[0][0],
        "qty": {str(k): v for k, v in p["qty"].most_common(4)}
    })
products.sort(key=lambda x: -x["n"]); products = products[:160]
seen_keys = set((p["name"], p["opt"]) for p in products)

# ---------- 단가비교: 최저가 구매처와 함께 (운송장에 없는 상품 보강) ----------
GRADE = re.compile(r"^(대과|소과|중과|로얄과|특대|상|중|하)\b")
def catalog_from(sheet):
    try: ws = wb[sheet]
    except KeyError: return []
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 3: return []
    vend = [str(v).strip() if v else "" for v in rows[1]]
    cur, out = None, []
    for r in rows[2:]:
        if r[1]: cur = re.sub(r"\s+", " ", str(r[1])).strip()
        if not cur or not r[2] or not isinstance(r[21], (int, float)) or r[21] <= 0: continue
        prices = {vend[i]: r[i] for i in range(27, min(len(r), 83)) if vend[i] and isinstance(r[i], (int, float)) and r[i] > 0}
        v = min(prices, key=prices.get) if prices else ""
        name = re.sub(r"\(.*?\)|\[.*?\]|/\s*[가-힣A-Za-z]+$", " ", cur)
        for vn in list(vendors.keys()) + ["업푸르트", "웰그린", "제이비티엠", "품담", "올바른푸드", "하루팜", "팡이농장", "덤덤몰", "최고집", "늘푸른우리", "팜허브"]:
            name = name.replace(vn, " ")
        name = tidy(re.sub(r"\s+", " ", name)).strip(" /")
        if len(name) < 2 or GRADE.match(name): continue
        opt = str(r[2]).strip(); opt = re.sub(r"^(\d+)\.0$", r"\1개", opt)
        low = r[3] if isinstance(r[3], (int, float)) and 0 < r[3] < r[21] else round(r[21] * 0.72, -1)
        out.append({"name": name[:24], "opt": opt[:24], "n": 40, "price": int(round(r[21], -1)), "cost": int(round(low, -1)), "vendor": norm_vendor(v) or "산지농가", "courier": "CJ대한통운", "qty": {"1": 9, "2": 1}})
    return out
for extra in catalog_from("단가비교"):
    k = (extra["name"], extra["opt"])
    if k in seen_keys: continue
    seen_keys.add(k); products.append(extra)

# ---------- 계절·명절 상품 보강 (실제 시세 수준의 가격, 구매처는 운영 중인 거래처 순환) ----------
VENDOR_CYCLE = [v for v, _ in vendors.most_common(12)] or ["성공푸드"]
def cur(name, opt, price, cost, n, courier="CJ대한통운", season=None, vendor=None):
    global _vi
    v = vendor or VENDOR_CYCLE[_vi % len(VENDOR_CYCLE)]; _vi += 1
    d = {"name": name, "opt": opt, "n": n, "price": price, "cost": cost, "vendor": v, "courier": courier, "qty": {"1": 9, "2": 1}}
    if season: d["season"] = season
    return d
_vi = 0
AUTUMN = {8: 0.6, 9: 1.6, 10: 1.4, 11: 1.0}
CURATED = [
    cur("햅쌀 (2026년산)", "10kg", 34900, 27500, 520, season={9: 1.4, 10: 1.6, 11: 1.3}),
    cur("햅쌀 (2026년산)", "20kg", 62900, 50800, 310, season={9: 1.4, 10: 1.6, 11: 1.3}),
    cur("신동진 쌀", "10kg", 32900, 26100, 260), cur("신동진 쌀", "20kg", 59900, 48500, 150),
    cur("현미", "10kg", 36900, 29400, 90), cur("찹쌀", "5kg", 24900, 18900, 70),
    cur("추석 선물세트 사과·배 혼합", "5kg (사과 5·배 4)", 49900, 36500, 240, season=AUTUMN),
    cur("추석 선물세트 나주배", "7.5kg (9~11과)", 59900, 44000, 180, season=AUTUMN),
    cur("문경 부사 사과", "3kg (12~15과)", 21900, 15400, 330, season=AUTUMN), cur("문경 부사 사과", "5kg (17~22과)", 32900, 23800, 210, season=AUTUMN),
    cur("홍로 사과", "2.5kg (가정용)", 16900, 11600, 260, season={8: 1.2, 9: 1.8, 10: 0.8}),
    cur("샤인머스캣", "2kg (2~3송이)", 26900, 19200, 420, season={8: 1.2, 9: 1.7, 10: 1.5, 11: 0.8}),
    cur("샤인머스캣", "4kg", 48900, 36000, 160, season={8: 1.2, 9: 1.7, 10: 1.5, 11: 0.8}),
    cur("캠벨 포도", "3kg", 19900, 13800, 230, season={8: 1.5, 9: 1.6, 10: 0.7}),
    cur("거봉", "2kg", 22900, 16100, 150, season={8: 1.4, 9: 1.5}),
    cur("무화과", "1kg (12~16과)", 18900, 12900, 170, season={8: 1.5, 9: 1.6, 10: 0.6}),
    cur("황도 복숭아", "3kg (10~13과)", 24900, 17200, 140, season={7: 1.6, 8: 1.8, 9: 0.9}),
    cur("공주 알밤", "2kg (특)", 19900, 13600, 210, season=AUTUMN), cur("건대추", "1kg", 21900, 15100, 120, season=AUTUMN),
    cur("성주 참외", "3kg", 19900, 13300, 60, season={5: 1.8, 6: 1.9, 7: 1.4, 8: 0.6}),
    cur("논산 설향 딸기", "1kg (2팩)", 21900, 15900, 60, season={12: 1.8, 1: 2.0, 2: 2.0, 3: 1.6, 4: 1.0}),
    cur("제주 한라봉", "3kg", 24900, 17500, 60, season={1: 1.8, 2: 1.9, 3: 1.4}),
    cur("골드키위", "1.5kg (12~15과)", 17900, 12800, 190), cur("블루베리", "500g", 15900, 11200, 110),
    cur("바나나", "1.3kg (1송이)", 5990, 4100, 380), cur("아보카도", "5개입", 9900, 7000, 160),
    cur("무항생제 계란", "30구 (대란)", 9990, 7300, 460), cur("무항생제 계란", "60구", 18900, 14000, 200),
    cur("햇양파", "5kg", 8990, 5900, 300), cur("깐마늘", "1kg", 11900, 8200, 210), cur("대파", "1kg", 4990, 3300, 180),
    cur("수미 감자", "5kg", 12900, 8700, 240), cur("밤호박 단호박", "3kg (3~4개)", 12900, 8800, 200, season={8: 1.4, 9: 1.5, 10: 1.2}),
    cur("초당 옥수수", "10개", 14900, 9800, 90, season={6: 1.8, 7: 1.9, 8: 1.3}),
    cur("애호박", "3개", 5490, 3600, 150), cur("오이", "10개", 7990, 5300, 130), cur("파프리카", "1kg", 8990, 6100, 110),
    cur("상추", "500g", 4990, 3200, 120), cur("깻잎", "200g", 3990, 2500, 100), cur("청양고추", "500g", 6990, 4600, 90),
    cur("양배추", "2통", 6990, 4500, 110), cur("당근", "3kg", 8990, 5900, 90), cur("브로콜리", "3개", 7990, 5200, 80),
    cur("가을 배추", "3포기", 14900, 9800, 130, season={10: 1.6, 11: 2.0}), cur("총각무", "2kg", 7990, 5100, 70, season={10: 1.4, 11: 1.6}),
    cur("표고버섯", "1kg", 15900, 11000, 120), cur("새송이버섯", "1kg", 6990, 4600, 100), cur("느타리버섯", "1kg", 6990, 4400, 80),
    cur("완도 재래김", "50봉 (도시락김)", 14900, 10200, 260), cur("국물용 멸치", "1.5kg", 24900, 17800, 90),
    cur("제주 은갈치", "4마리 (특대)", 39900, 29500, 70), cur("손질 고등어", "10팩", 27900, 20300, 110),
    cur("영광 굴비", "10미 (선물세트)", 59900, 44500, 60, season=AUTUMN),
    cur("횡성 한우 불고기", "1kg (1등급)", 42900, 33800, 90, season=AUTUMN), cur("돼지 삼겹살", "1kg", 21900, 16100, 130),
    cur("떡국 떡", "2kg", 12900, 9000, 60, season={1: 1.8, 2: 1.6}), cur("송편 (모둠)", "1kg", 14900, 10100, 120, season={9: 2.2, 10: 0.4}),
]
for extra in CURATED:
    k = (extra["name"], extra["opt"])
    if k in seen_keys: continue
    seen_keys.add(k); products.append(extra)
# 거래처가 빈 상품은 같은 상품명의 최다 거래처, 없으면 '산지농가'
by_name = collections.defaultdict(collections.Counter)
for p in products:
    if p["vendor"]: by_name[p["name"]][p["vendor"]] += p["n"]
for p in products:
    if not p["vendor"]: p["vendor"] = (by_name[p["name"]].most_common(1) or [("산지농가", 0)])[0][0]
if ANON:
    names = sorted(set(p["vendor"] for p in products), key=lambda v: -vendors.get(v, 0))
    alias = {v: "구매처 " + chr(65 + i) if i < 26 else "구매처 " + str(i + 1) for i, v in enumerate(names)}
    for p in products: p["vendor"] = alias[p["vendor"]]
    vendors = collections.Counter({alias.get(k, k): v for k, v in vendors.items() if k in alias})

# ---------- 일별 매출 흔들림 (매출요약) ----------
H2, it2 = rows_of("매출요약")
series = []
for r in it2:
    d = r[0]
    if isinstance(d, datetime.datetime) and num(r[3]): series.append((d.date(), float(r[3]), float(r[2] or 0)))
series.sort()
noise = []
for i, (d, rev, cnt) in enumerate(series):
    win = [x[1] for x in series[max(0, i - 3): i + 4]]
    m = sum(win) / len(win)
    if m > 0: noise.append(round(max(0.55, min(1.7, rev / m)), 3))
noise = noise[-120:]

# ---------- CS 유형 비율 (텍스트는 담지 않고 개수만) ----------
H3, it3 = rows_of("CS ")
cs = collections.Counter(); cs_n = 0
KW = [("반품접수", r"반품|회수|반송"), ("차액협의", r"차액|보상|부분환불|환급|입금"), ("출고중지요청", r"취소|출고중지|출고 중지"),
      ("택배미수령", r"미수령|분실|못 받|안 왔|안왔"), ("품절취소", r"품절"), ("과배송", r"과배송|오배송|잘못"), ("교환문의", r"교환"), ("배송문의", r"지연|언제|배송문의")]
for r in it3:
    if not r or len(r) < 8: continue
    txt = " ".join(str(x) for x in r[:8] if x)
    if not txt.strip() or not r[0]: continue
    cs_n += 1
    for k, pat in KW:
        if re.search(pat, txt): cs[k] += 1; break
    else: cs["기타CS"] += 1

data = {
    "source": "실제 운영 워크북 패턴(개인정보 제외) · 생성 " + datetime.date.today().isoformat(),
    "orders": n_rows,
    "products": products,
    "hours": hours, "weekdays": wdays,
    "qty": {str(k): v for k, v in qtys.most_common(6)},
    "regions": [[k, v] for k, v in regions.most_common(70)],
    "couriers": [[k, v] for k, v in couriers.most_common(8)],
    "vendors": [[k, v] for k, v in vendors.most_common(20)],
    "discounts": [[k, v] for k, v in discounts.most_common(5)],
    "noise": noise,
    "cs": {"total": cs_n, "mix": dict(cs.most_common())},
}
js = "/* demo-data.js — 자동화 관제 데모용 패턴 데이터. tools/make-demo-data.py 가 생성. 개인정보 없음(상품·가격·구매처·택배사·지역·시간대·CS 비율만) */\nwindow.DemoData = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n"
open(OUT, "w", encoding="utf-8").write(js)
print("rows", n_rows, "products", len(products), "regions", len(regions), "couriers", couriers.most_common(5), "vendors", vendors.most_common(6))
print("hours", hours); print("weekdays", wdays); print("qty", qtys.most_common(5)); print("discounts", discounts.most_common(5))
print("cs", cs_n, cs.most_common()); print("noise n", len(noise), noise[:10])
print("AOV", sum(p["price"] * p["n"] for p in products) / max(1, sum(p["n"] for p in products)))
for p in products[:25]: print("  ", p["n"], p["name"], "|", p["opt"], p["price"], p["cost"], p["vendor"], p["courier"])
