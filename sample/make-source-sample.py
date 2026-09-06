#!/usr/bin/env python3
"""합성 테스트 파일 생성 — 새 소스 양식(36열)과 같은 헤더에 가짜 주문을 채웁니다.

    python3 sample/make-source-sample.py            # sample/소스샘플_36열.xlsx / .xls / .csv 생성

개인정보가 전혀 들어가지 않는 가짜 데이터입니다. 이름·전화·주소·주문번호 모두 생성값.
"""
import csv
import datetime as dt
import os
import random

try:
    import openpyxl
except ImportError:
    openpyxl = None
try:
    import xlwt
except ImportError:
    xlwt = None

HEADER = ["수집일", "주문일", "판매사이트 주문번호", "판매사이트명", "판매자ID", "판매가", "배송비금액",
          "마스터상품코드", "판매사이트 상품코드", "상품명", "판매자상품코드", "주문선택사항", "주문수량", "에누리",
          "구매링크", "구매가", "구매자명", "수령자명", "수령자전화번호", "수령자휴대폰번호", "배송지우편번호",
          "배송지주소", "배송메세지", "담당자", "구매처", "계정", "구매금액", "주문번호　앞부분", "결제일시",
          "카드정보", "포인트", "주문여부", "한줄메모", "주문고유번호", "배송사명", "송장번호"]

random.seed(7)
SUR = "김이박최정강조윤장임한오서신권황안송류전홍"
GIVEN = ["민준", "서연", "도윤", "지우", "하준", "서윤", "예준", "하은", "지호", "수아", "시우", "지민", "유진", "현우", "은우", "채원"]
PRODUCTS = [
    ("수향미 골든퀸 3호 백미", "10kg 1개 특등급", 49800, 42100),
    ("모던하우스 스테인리스 물병 2P 세트", "크림 2개 1.5L", 20560, 15960),
    ("오리온 촉촉한 황치즈칩 16P", "320g x 4개", 23330, 18900),
    ("[지지피엑스] NEW 아트웍 폰테 티셔츠", "사이즈:55", 38600, 29670),
    ("아기물티슈 캡형 100매", "10팩", 15900, 11200),
    ("세탁세제 액체 3L", "2개", 21900, 16800),
]
CITIES = [
    ("서울특별시 마포구 월드컵북로 {n} {b}동 {h}호", "0{z}"),
    ("경기도 용인시 기흥구 흥덕1로 {n} {b}동 {h}호", "1{z}"),
    ("부산광역시 해운대구 센텀중앙로 {n} {b}동 {h}호", "4{z}"),
    ("인천광역시 연수구 송도과학로 {n} {b}동 {h}호", "2{z}"),
    ("대구광역시 수성구 동대구로 {n} {b}동 {h}호", "4{z}"),
]
VENDORS = ["11번가", "지마켓", "옥션", "롯데온"]
CARDS = ["롯데", "삼성", "하나"]


def person():
    return random.choice(SUR) + random.choice(GIVEN)


def phone():
    return "0502-%04d-%04d" % (random.randint(1000, 9999), random.randint(1000, 9999))


def rows(n=24):
    out = []
    base = dt.datetime(2026, 9, 5, 9, 0)
    for i in range(n):
        p = random.choice(PRODUCTS)
        city, zp = random.choice(CITIES)
        ordered = base + dt.timedelta(minutes=random.randint(0, 60 * 30))
        collected = ordered + dt.timedelta(hours=random.randint(1, 20))
        site_no = "%013d %015d" % (random.randint(10 ** 12, 10 ** 13 - 1), random.randint(10 ** 14, 10 ** 15 - 1))
        done = i % 3 == 0   # 3건 중 1건은 이미 발주된 상태
        name = person()
        tel = phone()
        r = {
            "수집일": collected, "주문일": ordered, "판매사이트 주문번호": site_no, "판매사이트명": "쿠팡(신)",
            "판매자ID": "samplestore", "판매가": p[2], "배송비금액": 0,
            "마스터상품코드": "", "판매사이트 상품코드": str(random.randint(7 * 10 ** 9, 9 * 10 ** 9)),
            "상품명": p[0], "판매자상품코드": "", "주문선택사항": p[1], "주문수량": random.choice([1, 1, 1, 2]),
            "에누리": 0, "구매링크": "", "구매가": "",
            "구매자명": name, "수령자명": name if random.random() < 0.7 else person(),
            "수령자전화번호": tel, "수령자휴대폰번호": tel,
            "배송지우편번호": zp.format(z="%04d" % random.randint(1000, 9999)),
            "배송지주소": city.format(n=random.randint(10, 300), b=random.randint(101, 115), h=random.randint(101, 1504)),
            "배송메세지": random.choice(["문 앞", "경비실", "직접 받고 부재 시 문 앞", ""]),
            "담당자": "명진" if done else "", "구매처": random.choice(VENDORS) if done else "",
            "계정": "", "구매금액": p[3] if done else "",
            "주문번호　앞부분": str(random.randint(10 ** 9, 10 ** 10 - 1)) if done else "",
            "결제일시": collected + dt.timedelta(hours=2) if done else "",
            "카드정보": random.choice(CARDS) if done else "", "포인트": random.randint(100, 500) if done else "",
            "주문여부": "O" if done else "", "한줄메모": "",
            "주문고유번호": str(random.randint(10 ** 11, 10 ** 12 - 1)),
            "배송사명": random.choice(["T025", "T081"]) if done else "",
            "송장번호": str(random.randint(10 ** 11, 10 ** 12 - 1)) if done else "",
        }
        if done:
            r["계정"] = r["구매처"] + "(명진)"
        out.append(r)
    return out


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    data = rows()
    # CSV
    with open(os.path.join(here, "소스샘플_36열.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        for r in data:
            w.writerow([r[h].strftime("%Y-%m-%d %H:%M") if isinstance(r[h], dt.datetime) else r[h] for h in HEADER])
    # XLSX
    if openpyxl:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "2026-09-07"
        ws.append(HEADER)
        for r in data:
            ws.append([r[h] for h in HEADER])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, dt.datetime):
                    c.number_format = "yyyy-mm-dd hh:mm"
        wb.save(os.path.join(here, "소스샘플_36열.xlsx"))
    # XLS (BIFF8)
    if xlwt:
        wb2 = xlwt.Workbook(encoding="utf-8")
        sh = wb2.add_sheet("2026-09-07")
        datefmt = xlwt.easyxf(num_format_str="yyyy-mm-dd hh:mm")
        for c, h in enumerate(HEADER):
            sh.write(0, c, h)
        for ri, r in enumerate(data, start=1):
            for c, h in enumerate(HEADER):
                v = r[h]
                if isinstance(v, dt.datetime):
                    sh.write(ri, c, v, datefmt)
                elif v == "":
                    continue
                else:
                    sh.write(ri, c, v)
        wb2.save(os.path.join(here, "소스샘플_36열.xls"))
    # 헤더만 있는 빈 양식(수집 프로그램이 준 샘플과 같은 형태)
    if xlwt:
        wb3 = xlwt.Workbook(encoding="utf-8")
        sh3 = wb3.add_sheet("2026-09-07")
        for c, h in enumerate(HEADER):
            sh3.write(0, c, h)
        wb3.save(os.path.join(here, "소스양식_빈파일.xls"))
    print("생성:", len(data), "행 →", here)


if __name__ == "__main__":
    main()
